"""DOCX 正文/表格及 Markdown/TXT 入库；截图 OCR 尚未实现。"""

from dataclasses import dataclass
from hashlib import sha256
import json
from pathlib import Path
import re
from uuid import NAMESPACE_URL, uuid5
from urllib.parse import parse_qsl, urlsplit, urlunsplit
from zipfile import ZipFile

SUPPORTED_SUFFIXES = {".docx", ".md", ".txt"}


@dataclass
class ParsedDocument:
    id: str
    path: Path
    blocks: list[tuple[str, str]]
    metadata: dict
    images_unprocessed: int = 0


def redact(text: str) -> str:
    """清理日志和正文中的凭据，同时保留相邻错误信息及普通文档 URL。"""
    text = re.sub(
        r"-----BEGIN [^-]*PRIVATE KEY-----.*?-----END [^-]*PRIVATE KEY-----",
        "[REDACTED]",
        text,
        flags=re.S,
    )

    def clean_url(match):
        original = match.group(0)
        # 句号、Markdown 闭括号等属于正文，清理 URL 后再放回。
        value = original.rstrip(".,;!?)]}，。；")
        suffix = original[len(value):]
        try:
            parsed = urlsplit(value)
            netloc = parsed.netloc.rsplit("@", 1)[-1]
            sensitive = re.compile(
                r"(?i)(password|passwd|token|secret|api[_-]?key|authorization|"
                r"signature|credential|access[_-]?key|awsaccesskeyid|(?:^|[_-])sig$)"
            )
            signed = any(sensitive.search(key) for key, _ in parse_qsl(parsed.query))
            fragment_private = any(
                sensitive.search(key) for key, _ in parse_qsl(parsed.fragment)
            )
            return urlunsplit((parsed.scheme, netloc, parsed.path,
                               "" if signed else parsed.query,
                               "" if fragment_private or signed else parsed.fragment)) + suffix
        except ValueError:
            return "[REDACTED_URL]" + suffix

    text = re.sub(
        r'''(?i)\b(?:https?|postgres(?:ql)?|mysql|redis|amqp|ssh)://[^\s<>"'`]+''',
        clean_url, text,
    )
    # 引号形式同时适配 JSON/YAML；只替换字段值，保留同一行的故障上下文。
    labels = (
        r"password|passwd|pwd|token|access[_-]?token|refresh[_-]?token|secret|"
        r"secret[_-]?key|access[_-]?key(?:[_-]?id|[_-]?secret)?|"
        r"client[_-]?secret|client[_-]?key[_-]?data|api[_-]?key|authorization|private[_-]?key"
    )
    text = re.sub(
        rf'''(?ix)((?<![\w])["']?(?:{labels})["']?[ \t]*[:=：][ \t]*)(?:"(?:\\.|[^"\\\n])*"|'(?:\\.|[^'\\\n])*'|(?:Bearer|Basic)[ \t]+[^\s,;}}]+|[^\s,;}}]+)''',
        r"\1[REDACTED]", text,
    )
    text = re.sub(
        r'''((?:密码|口令)["']?[ \t]*[:：=][ \t]*)(?:"(?:\\.|[^"\\\n])*"|'(?:\\.|[^'\\\n])*'|[^\s,;}]+)''',
        r"\1[REDACTED]", text,
    )
    text = re.sub(r"(?i)\bBearer[ \t]+[A-Za-z0-9._~+/-]+=*", "Bearer [REDACTED]", text)
    text = re.sub(r"sk-[A-Za-z0-9_-]{16,}", "[REDACTED]", text)
    return text


def source_paths(directory: Path) -> list[Path]:
    root = directory.resolve()
    return [
        path
        for path in sorted(directory.rglob("*"))
        if path.is_file()
        and not path.is_symlink()
        and path.suffix.lower() in SUPPORTED_SUFFIXES
        and path.resolve().is_relative_to(root)
    ]


# ParsedDocument(
#     id="文档ID",
#     path=Path(...),
#     blocks=[
#         ("安装准备 · 行 3", "准备离线安装包……"),
#         ("导入镜像 · 行 8", "执行镜像导入……"),
#     ],
#     metadata={...},
#     images_unprocessed=0,
# )
def read_document(path: Path, root: Path) -> ParsedDocument:
    relative = path.relative_to(root).as_posix()
    document_id = sha256(relative.encode()).hexdigest()[:16]
    metadata = {
        "source": relative,
        "title": path.stem,
        "kind": "document",
        "kubernetes_versions": ["any"],
        "operating_systems": ["any"],
        "install_modes": ["online", "offline"],
        "verification_status": "source_document",
    }
    sidecar = path.with_suffix(path.suffix + ".metadata.json")
    if sidecar.exists():
        additions = json.loads(sidecar.read_text(encoding="utf-8"))
        for key in (
            "title",
            "kind",
            "kubernetes_versions",
            "operating_systems",
            "install_modes",
            "verification_status",
        ):
            if key in additions:
                metadata[key] = additions[key]
    for key in ("kubernetes_versions", "operating_systems", "install_modes"):
        if (
            not isinstance(metadata[key], list)
            or not metadata[key]
            or not all(isinstance(v, str) for v in metadata[key])
        ):
            raise ValueError(f"{path.name} 的 {key} 必须是非空字符串列表。")
    blocks = []
    images = 0
    if path.suffix.lower() == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(path)
        section = "正文"
        for number, block in enumerate(document.iter_inner_content(), 1):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if (
                    text
                    and block.style
                    and block.style.name.lower().startswith(("heading", "标题"))
                ):
                    section = text
            elif isinstance(block, Table):
                text = "\n".join(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in block.rows
                )
            else:
                continue
            if text:
                blocks.append((f"{section} · 内容块 {number}", redact(text)))
        with ZipFile(path) as archive:
            images = sum(
                name.startswith("word/media/") and not name.endswith("/")
                for name in archive.namelist()
            )
    else:
        section = "正文"
        for number, text in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if text.startswith("#"):
                section = text.lstrip("#").strip() or section
            if text.strip():
                blocks.append((f"{section} · 行 {number}", redact(text.strip())))
    return ParsedDocument(document_id, path, blocks, metadata, images)


def chunk_document(document: ParsedDocument, size: int, overlap: int) -> list[dict]:
    # 按内容块保留定位，长块按固定长度切分；重复正文去重但保留首个来源定位。
    output = []
    seen = set()
    buffer = ""
    locations = []

    def append(content: str, location: str) -> None:
        normalized = content.strip()
        if not normalized or normalized in seen:
            return
        seen.add(normalized)
        metadata = {**document.metadata, "location": location}
        identity = (
            json.dumps(metadata, ensure_ascii=False, sort_keys=True) + "\n" + normalized
        )
        output.append(
            {
                "id": str(
                    uuid5(NAMESPACE_URL, f"private-rag:{document.id}:{identity}")
                ),
                "document_id": document.id,
                "content": normalized,
                "metadata": metadata,
            }
        )

    for location, text in document.blocks:
        if buffer and len(buffer) + len(text) + 1 > size:
            append(buffer, "；".join(dict.fromkeys(locations)))
            buffer = ""
            locations = []
        if len(text) > size:
            if buffer:
                append(buffer, "；".join(dict.fromkeys(locations)))
                buffer = ""
                locations = []
            for start in range(0, len(text), size - overlap):
                append(text[start : start + size], location)
                if start + size >= len(text):
                    break
        else:
            buffer += ("\n" if buffer else "") + text
            locations.append(location)
    append(buffer, "；".join(dict.fromkeys(locations)))
    return output
